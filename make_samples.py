# make_samples.py  (project root)

from pathlib import Path
try:
    import pymupdf
except ImportError:
    import fitz as pymupdf
from docx import Document

OUT = Path("sample-docs")
OUT.mkdir(exist_ok=True)
 
FAKE = (
    "Dear Mr. Fred Banda, thank you for your FOI request. "
    "We have your details on file: date of birth 12/03/1985, "
    "PPS number 1234567AB, phone 087 123 4567, "
    "email fred.banda@example.com, "
    "address 14 Oak Road, Santry, Dublin 9."
)
 # 1. plain text - used by the Day 3 evaluation script
(OUT / "sample1.txt").write_text(FAKE, encoding="utf-8")
 

# 2.Word sample
d = Document()
d.add_heading("FOI Response – Internal Draft", level=1)
d.add_paragraph(FAKE)
d.save("sample_docs/sample1.docx")

# 3. Word with a table, a header and a footer (the usual leak paths)
d2 = Document()
d2.sections[0].header.paragraphs[0].text = (
    "Case officer: Aoife Ni Bhriain, aoife.nibhriain@example.gov.ie")
d2.sections[0].footer.paragraphs[0].text = "Reviewed by Sean O Conghaile"
d2.add_heading("Schedule of Records", level=1)
t = d2.add_table(rows=2, cols=2)
t.cell(0, 0).text = "Name";              t.cell(0, 1).text = "PPS number"
t.cell(1, 0).text = "Sean O Conghaile";  t.cell(1, 1).text = "9876543CD"
d2.save(OUT / "sample2.docx")
 

# 4. Email
(OUT / "sample1.eml").write_text(
    "From: Mary O'Brien <mary.obrien@example.com>\n"
    "To: FOI Unit <foi@example.gov.ie>\n"
    "Subject: Request for records\n\n" + FAKE,
    encoding="utf-8")
 
# 5. PDF, generated so you never have to open Word
pdf = pymupdf.open()
page = pdf.new_page()
page.insert_textbox(pymupdf.Rect(60, 60, 540, 720),
                    "FOI Response - Internal Draft\n\n" + FAKE,
                    fontsize=11, fontname="helv")
pdf.save(OUT / "sample1.pdf")
pdf.close()

print("Samples created in", OUT.resolve())

