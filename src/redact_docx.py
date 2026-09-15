# src/redact_docx.py
from docx import Document
from .detector import PIIDetector, redact_text
 
 
def _cell_paragraphs(container):
    paras = []
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                paras.extend(cell.paragraphs)
    return paras
 
 
def _all_paragraphs(doc):
    """Body, tables, headers and footers. Headers are a common leak path:
    the case officer's name sits there on every page."""
    paras = list(doc.paragraphs) + _cell_paragraphs(doc)
    for section in doc.sections:
        for part in (section.header, section.footer,
                     section.first_page_header, section.first_page_footer):
            paras.extend(part.paragraphs)
            paras.extend(_cell_paragraphs(part))
    return paras
 
 
def redact_docx(in_path: str, out_path: str, det: PIIDetector):
    doc = Document(in_path)
    all_findings = []
    for i, para in enumerate(_all_paragraphs(doc)):
        if not para.text.strip():
            continue
        findings = det.analyze(para.text)
        if not findings:
            continue
        new_text = redact_text(para.text, findings)
        for run in para.runs:          # clear the existing runs
            run.text = ""
        if para.runs:
            para.runs[0].text = new_text
        else:
            para.add_run(new_text)
        all_findings.extend((f, f"paragraph {i}") for f in findings)
    doc.save(out_path)
    return all_findings
