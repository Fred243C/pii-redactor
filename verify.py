# verify.py  (project root)
"""Fail-loud leak test: re-read every redacted output and search for the
known fake PII. Exit code 1 if anything survived. This is the single most
useful script in the project - run it after every change."""
import sys
from pathlib import Path
try:
    import pymupdf                 # PyMuPDF 1.24.3 and later
except ImportError:
    import fitz as pymupdf         # older releases expose it as fitz
from docx import Document
 
NEEDLES = [
    "John Murphy", "1234567AB", "087 123 4567", "john.murphy@example.com",
    "14 Oak Road", "Santry", "Mary O'Brien", "mary.obrien@example.com",
    "Aoife Ni Bhriain", "Sean O Conghaile", "9876543CD",
]
 
 
def text_of(path: Path) -> str:
    if path.suffix == ".pdf":
        with pymupdf.open(path) as doc:
            return "\n".join(page.get_text() for page in doc)
    if path.suffix == ".docx":
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                parts.extend(c.text for c in row.cells)
        for s in doc.sections:
            parts.extend(p.text for p in s.header.paragraphs)
            parts.extend(p.text for p in s.footer.paragraphs)
        return "\n".join(parts)
    return path.read_text(encoding="utf-8", errors="ignore")
 
 
leaked = False
for out in sorted(Path("out").glob("*_redacted.*")):
    body = text_of(out).lower()
    hits = [n for n in NEEDLES if n.lower() in body]
    if hits:
        leaked = True
        print(f"LEAK  {out.name}: {hits}")
    else:
        print(f"CLEAN {out.name}")
sys.exit(1 if leaked else 0)
